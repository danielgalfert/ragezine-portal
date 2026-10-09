from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core import mail
from django.test import TestCase
from django.test import override_settings
from django.conf import settings
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient
from io import BytesIO
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from types import SimpleNamespace
from unittest.mock import patch

from submissions.api.downloads import _build_complete_download_zip
from submissions.api.email import send_submission_receipt
from submissions.models import Submission, SubmissionDocument
from submissions.repositories import SubmissionDocumentRepository, SubmissionRepository


class AuthenticationEndpointsTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = get_user_model().objects.create_user(
            username="editor",
            password="S3cretPass123",
            is_staff=True,
        )

    def test_login_session_and_logout_flow(self):
        login_response = self.client.post(
            "/api/auth/login/",
            {"username": "editor", "password": "S3cretPass123"},
            format="json",
        )

        self.assertEqual(login_response.status_code, status.HTTP_200_OK)
        self.assertEqual(login_response.data["user"]["username"], self.user.username)

        session_response = self.client.get("/api/auth/session/")
        self.assertEqual(session_response.status_code, status.HTTP_200_OK)
        self.assertTrue(session_response.data["authenticated"])

        logout_response = self.client.post("/api/auth/logout/")
        self.assertEqual(logout_response.status_code, status.HTTP_200_OK)

        session_after_logout = self.client.get("/api/auth/session/")
        self.assertEqual(session_after_logout.status_code, status.HTTP_200_OK)
        self.assertFalse(session_after_logout.data["authenticated"])

    def test_login_rejects_invalid_credentials(self):
        response = self.client.post(
            "/api/auth/login/",
            {"username": "editor", "password": "wrong"},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data["message"], "Invalid username or password.")


class SubmissionValidationTests(TestCase):
    def setUp(self):
        self.client = APIClient()

    def base_payload(self):
        return {
            "title": "A Valid Title",
            "submission_type": "poetry",
            "artist_name": "Artist Name",
            "pronouns": "they/them",
            "short_bio": "A short but valid bio for the submission.",
            "country_origin": "DK",
            "countries_residence": ["DK"],
            "language": "English",
            "allow_translation": True,
            "text_files": SimpleUploadedFile("poem.pdf", b"test file content", content_type="application/pdf"),
        }

    def test_submission_requires_email(self):
        response = self.client.post("/api/submissions/", self.base_payload(), format="multipart")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("email", response.data)

    def test_submission_rejects_invalid_email(self):
        payload = self.base_payload()
        payload["email"] = "not-an-email"

        response = self.client.post("/api/submissions/", payload, format="multipart")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("email", response.data)

    def test_non_english_submission_without_translation_checkbox(self):
        payload = self.base_payload()
        payload["email"] = "artist@example.com"
        payload["language"] = "Danish"
        payload.pop("allow_translation")

        response = self.client.post("/api/submissions/", payload, format="multipart")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        submission = Submission.objects.get(pk=response.data["id"])
        self.assertFalse(submission.allow_translation)

    def test_submission_uploads_are_stored_as_documents(self):
        payload = self.base_payload()
        payload["email"] = "artist@example.com"
        payload["visuals"] = SimpleUploadedFile(
            "image.tiff",
            b"image content",
            content_type="image/tiff",
        )

        response = self.client.post("/api/submissions/", payload, format="multipart")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        submission = Submission.objects.get(pk=response.data["id"])
        documents = submission.documents.order_by("document_type", "original_filename")

        self.assertEqual(documents.count(), 2)
        self.assertEqual(
            [document.document_type for document in documents],
            [
                SubmissionDocument.DocumentType.TEXT,
                SubmissionDocument.DocumentType.VISUAL,
            ],
        )
        self.assertEqual(len(response.data["documents"]), 2)
        self.assertTrue(all("file" not in item for item in response.data["documents"]))
        self.assertTrue(all("/api/downloads/documents/" in item["download_url"] for item in response.data["documents"]))

        staff_user = get_user_model().objects.create_user(
            username="document-editor", password="S3cretPass123", is_staff=True
        )
        self.client.force_authenticate(staff_user)
        text_document = documents.get(document_type=SubmissionDocument.DocumentType.TEXT)
        download_response = self.client.get(f"/api/downloads/documents/{text_document.pk}/")
        self.assertEqual(download_response.status_code, status.HTTP_200_OK)
        self.assertEqual(b"".join(download_response.streaming_content), b"test file content")
        self.assertTrue(download_response.closed)
        self.assertEqual(self.client.get(f"/media/{text_document.file.name}").status_code, 404)

    def test_video_is_stored_as_visual_document(self):
        payload = self.base_payload()
        payload["email"] = "artist@example.com"
        payload["visuals"] = SimpleUploadedFile(
            "clip.mp4", b"video bytes", content_type="video/mp4"
        )

        with patch("submissions.api.submission_views.send_submission_receipt"):
            response = self.client.post(
                "/api/submissions/", payload, format="multipart", REMOTE_ADDR="192.0.2.53"
            )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        video = SubmissionDocument.objects.get(original_filename="clip.mp4")
        self.assertEqual(video.document_type, SubmissionDocument.DocumentType.VISUAL)
        self.assertEqual(video.size, len(b"video bytes"))

        staff_user = get_user_model().objects.create_user(
            username="video-editor", password="S3cretPass123", is_staff=True
        )
        self.client.force_authenticate(staff_user)
        download = self.client.get(f"/api/downloads/documents/{video.pk}/")
        self.assertEqual(download.status_code, status.HTTP_200_OK)
        self.assertEqual(b"".join(download.streaming_content), b"video bytes")
        self.assertTrue(download.closed)


class SubmissionSizeLimitTests(TestCase):
    def setUp(self):
        self.client = APIClient()

    def base_payload(self):
        return {
            "title": "A Valid Title",
            "submission_type": "poetry",
            "artist_name": "Artist Name",
            "pronouns": "they/them",
            "short_bio": "A short but valid bio for the submission.",
            "country_origin": "DK",
            "countries_residence": ["DK"],
            "language": "English",
            "allow_translation": True,
            "email": "artist@example.com",
        }

    @override_settings(
        SUBMISSION_MAX_FILE_BYTES=1024 * 1024,
        SUBMISSION_MAX_TOTAL_FILES_BYTES=2 * 1024 * 1024,
    )
    def test_rejects_file_over_individual_limit_without_saving_submission(self):
        payload = self.base_payload()
        payload["text_files"] = SimpleUploadedFile("large.pdf", b"x" * (1024 * 1024 + 1))

        response = self.client.post(
            "/api/submissions/", payload, format="multipart", REMOTE_ADDR="192.0.2.50"
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("text_files", response.data)
        self.assertFalse(Submission.objects.exists())

    @override_settings(
        SUBMISSION_MAX_FILE_BYTES=1024 * 1024,
        SUBMISSION_MAX_TOTAL_FILES_BYTES=1024 * 1024,
    )
    def test_rejects_files_over_combined_limit_without_saving_submission(self):
        payload = self.base_payload()
        payload["text_files"] = SimpleUploadedFile("poem.pdf", b"x" * 600_000)
        payload["visuals"] = SimpleUploadedFile("image.tiff", b"x" * 600_000)

        response = self.client.post(
            "/api/submissions/", payload, format="multipart", REMOTE_ADDR="192.0.2.51"
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("non_field_errors", response.data)
        self.assertFalse(Submission.objects.exists())

    @override_settings(
        SUBMISSION_MAX_FILE_BYTES=1024 * 1024,
        SUBMISSION_MAX_TOTAL_FILES_BYTES=2 * 1024 * 1024,
    )
    def test_accepts_files_at_both_limits(self):
        payload = self.base_payload()
        payload["text_files"] = SimpleUploadedFile("poem.pdf", b"x" * (1024 * 1024))
        payload["visuals"] = SimpleUploadedFile("image.tiff", b"x" * (1024 * 1024))

        with patch("submissions.api.submission_views.send_submission_receipt"):
            response = self.client.post(
                "/api/submissions/", payload, format="multipart", REMOTE_ADDR="192.0.2.52"
            )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(SubmissionDocument.objects.count(), 2)


class SubmissionRepositoryTests(TestCase):
    def test_repositories_create_submission_and_documents(self):
        submission_repository = SubmissionRepository()
        document_repository = SubmissionDocumentRepository()

        submission = submission_repository.create(
            {
                "title": "Repository Piece",
                "description": "Detailed description for repository testing.",
                "submission_type": Submission.SubmissionType.POETRY,
                "artist_name": "Repository Artist",
                "pronouns": "they/them",
                "short_bio": "A short but valid bio for repository testing.",
                "email": "repo@example.com",
                "country_origin": "DK",
                "countries_residence": ["DK"],
                "language": "English",
                "allow_translation": True,
            }
        )
        document_repository.create_text_document(
            submission=submission,
            uploaded_file=SimpleUploadedFile(
                "repo-piece.pdf",
                b"repo content",
                content_type="application/pdf",
            ),
        )

        loaded = submission_repository.get_with_documents(submission.id)

        self.assertEqual(loaded.title, "Repository Piece")
        self.assertEqual(loaded.documents.count(), 1)
        self.assertEqual(
            loaded.documents.first().document_type,
            SubmissionDocument.DocumentType.TEXT,
        )

    def test_same_named_uploads_get_distinct_storage_keys(self):
        submission = Submission.objects.create(
            title="Same Name",
            submission_type=Submission.SubmissionType.POETRY,
            artist_name="Artist",
            pronouns="they/them",
            short_bio="Bio",
            email="artist@example.com",
            country_origin="DK",
            countries_residence=["DK"],
        )
        repository = SubmissionDocumentRepository()
        first = repository.create_text_document(
            submission, SimpleUploadedFile("poem.pdf", b"first")
        )
        second = repository.create_text_document(
            submission, SimpleUploadedFile("poem.pdf", b"second")
        )

        self.assertNotEqual(first.file.name, second.file.name)
        self.assertTrue(first.file.name.startswith(f"submissions/texts/{submission.pk}/"))
        self.assertEqual(first.original_filename, "poem.pdf")
        self.assertEqual(second.original_filename, "poem.pdf")


class SubmissionDeletionTests(TestCase):
    def setUp(self):
        media_dir = TemporaryDirectory()
        self.addCleanup(media_dir.cleanup)
        storage_override = override_settings(
            STORAGES={
                **settings.STORAGES,
                "default": {
                    "BACKEND": "django.core.files.storage.FileSystemStorage",
                    "OPTIONS": {"location": media_dir.name},
                },
            }
        )
        storage_override.enable()
        self.addCleanup(storage_override.disable)
        self.client = APIClient()
        self.client.force_authenticate(
            get_user_model().objects.create_user(
                username="delete-editor", password="S3cretPass123", is_staff=True
            )
        )

    def create_submission(self):
        return Submission.objects.create(
            title="Delete Piece",
            submission_type=Submission.SubmissionType.POETRY,
            artist_name="Artist",
            pronouns="they/them",
            short_bio="Bio",
            email="artist@example.com",
            country_origin="DK",
            countries_residence=["DK"],
        )

    def test_delete_removes_submission_and_its_file(self):
        submission = self.create_submission()
        document = SubmissionDocument.create_from_upload(
            submission,
            SimpleUploadedFile("poem.pdf", b"poem"),
            SubmissionDocument.DocumentType.TEXT,
        )
        storage = document.file.storage
        name = document.file.name

        response = self.client.delete(f"/api/submissions/{submission.pk}/")

        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Submission.objects.filter(pk=submission.pk).exists())
        self.assertFalse(storage.exists(name))

    def test_delete_preserves_file_referenced_by_another_submission(self):
        first = self.create_submission()
        document = SubmissionDocument.create_from_upload(
            first,
            SimpleUploadedFile("shared.pdf", b"shared"),
            SubmissionDocument.DocumentType.TEXT,
        )
        second = self.create_submission()
        SubmissionDocument.objects.create(
            submission=second,
            document_type=SubmissionDocument.DocumentType.TEXT,
            file=document.file.name,
            original_filename="shared.pdf",
        )
        storage = document.file.storage
        name = document.file.name

        response = self.client.delete(f"/api/submissions/{first.pk}/")

        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertTrue(storage.exists(name))


class SubmissionDownloadsTests(TestCase):
    def setUp(self):
        media_dir = TemporaryDirectory()
        self.addCleanup(media_dir.cleanup)
        storage_override = override_settings(
            STORAGES={
                **settings.STORAGES,
                "default": {
                    "BACKEND": "django.core.files.storage.FileSystemStorage",
                    "OPTIONS": {"location": media_dir.name},
                },
            }
        )
        storage_override.enable()
        self.addCleanup(storage_override.disable)
        self.client = APIClient()
        self.user = get_user_model().objects.create_user(
            username="admin",
            password="S3cretPass123",
            is_staff=True,
        )
        self.client.force_authenticate(self.user)

    def test_complete_download_groups_submissions_and_files(self):
        submission = Submission.objects.create(
            title="Archive Piece",
            year="2026",
            description="Detailed description for export output.",
            submission_type=Submission.SubmissionType.POETRY,
            artist_name="Artist One",
            pronouns="they/them",
            short_bio="A short but valid bio for export testing.",
            socials="@artist",
            email="artist@example.com",
            country_origin="DK",
            countries_residence=["DK"],
            language="English",
            allow_translation=True,
        )
        SubmissionDocument.create_from_upload(
            submission=submission,
            uploaded_file=SimpleUploadedFile("legacy.txt", b"legacy file"),
            document_type=SubmissionDocument.DocumentType.TEXT,
        )
        SubmissionDocument.create_from_upload(
            submission=submission,
            uploaded_file=SimpleUploadedFile("poem.pdf", b"poem content", content_type="application/pdf"),
            document_type=SubmissionDocument.DocumentType.TEXT,
        )
        SubmissionDocument.create_from_upload(
            submission=submission,
            uploaded_file=SimpleUploadedFile("image.tiff", b"image content", content_type="image/tiff"),
            document_type=SubmissionDocument.DocumentType.VISUAL,
        )

        response = self.client.get("/api/downloads/complete/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "application/zip")

        with ZipFile(BytesIO(b"".join(response.streaming_content))) as archive:
            names = archive.namelist()
            self.assertIn("poetry/", names)
            root = f"poetry/Artist_One-{submission.pk}"
            self.assertIn(f"{root}/submission_details.txt", names)
            self.assertIn(f"{root}/legacy.txt", names)
            self.assertIn(f"{root}/poem.pdf", names)
            self.assertIn(f"{root}/image.tiff", names)

            details = archive.read(f"{root}/submission_details.txt").decode("utf-8")
            self.assertIn("Title: Archive Piece", details)
            self.assertIn("Email: artist@example.com", details)
            self.assertIn("- legacy.txt", details)
            self.assertIn("- poem.pdf", details)
            self.assertIn("- image.tiff", details)
        self.assertTrue(response.closed)

    def test_complete_download_separates_submissions_with_same_artist_name(self):
        submissions = []
        for content in (b"first", b"second"):
            submission = Submission.objects.create(
                title="Same Name",
                submission_type=Submission.SubmissionType.POETRY,
                artist_name="Artist One",
                pronouns="they/them",
                short_bio="Bio",
                email="artist@example.com",
                country_origin="DK",
                countries_residence=["DK"],
            )
            SubmissionDocument.create_from_upload(
                submission,
                SimpleUploadedFile("poem.pdf", content),
                SubmissionDocument.DocumentType.TEXT,
            )
            submissions.append(submission)

        with _build_complete_download_zip() as complete_zip, ZipFile(complete_zip) as archive:
            for submission, content in zip(submissions, (b"first", b"second")):
                self.assertEqual(
                    archive.read(f"poetry/Artist_One-{submission.pk}/poem.pdf"),
                    content,
                )

    def test_excel_download_contains_review_sheet_and_submission_row(self):
        Submission.objects.create(
            title="Spreadsheet Piece",
            year="2026",
            description="Detailed description for spreadsheet output.",
            submission_type=Submission.SubmissionType.LONG_FORM_WRITING,
            artist_name="Artist Two",
            pronouns="she/her",
            short_bio="A short but valid bio for spreadsheet export.",
            socials="@artist-two",
            email="artist.two@example.com",
            country_origin="DK",
            countries_residence=["DE"],
            language="English",
            allow_translation=False,
        )

        response = self.client.get("/api/downloads/excel/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('filename="rage_submissions_volume3.xlsx"', response["Content-Disposition"])

        workbook = load_workbook(filename=BytesIO(response.content))
        self.assertIn("Poetry", workbook.sheetnames)
        self.assertIn("Long form writing", workbook.sheetnames)

        sheet = workbook["Long form writing"]
        header_row = [cell.value for cell in sheet[1]]
        submission_row = [cell.value for cell in sheet[2]]
        expected_headers = [
            "ID",
            "Title",
            "Year",
            "Description",
            "Submission Type",
            "Artist Name",
            "Pronouns",
            "Short Bio",
            "Socials",
            "Email",
            "Country Origin",
            "Countries Residence",
            "Language",
            "Created At",
        ]

        self.assertEqual(header_row, expected_headers)
        self.assertEqual(submission_row[1], "Spreadsheet Piece")
        self.assertEqual(submission_row[2], "2026")
        self.assertEqual(submission_row[4], "Long form writing")
        self.assertEqual(submission_row[5], "Artist Two")
        self.assertEqual(submission_row[9], "artist.two@example.com")
        self.assertEqual(submission_row[10], "Denmark")
        self.assertEqual(submission_row[11], "Germany")


class SubmissionEmailTests(TestCase):
    @override_settings(
        EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
        SUBMISSION_EMAIL_FROM="editor@ragezine.test",
        SUBMISSION_EMAIL_SUBJECT="Thank you for submitting to Rage Zine",
        SUBMISSION_EMAIL_BODY="Thank you for submitting to Rage Zine.\n\nWe will review your work shortly.",
    )
    def test_send_submission_receipt_only_to_artist_without_attachments(self):
        submission = Submission.objects.create(
            title="Mail Piece",
            year="2026",
            description="Detailed description for email output.",
            submission_type=Submission.SubmissionType.POETRY,
            artist_name="Artist Mail",
            pronouns="they/them",
            short_bio="A short but valid bio for email testing.",
            socials="@artistmail",
            email="artist.mail@example.com",
            country_origin="DK",
            countries_residence=["DE"],
            language="English",
            allow_translation=True,
        )
        SubmissionDocument.create_from_upload(
            submission=submission,
            uploaded_file=SimpleUploadedFile("mail-piece.pdf", b"primary text"),
            document_type=SubmissionDocument.DocumentType.TEXT,
        )
        SubmissionDocument.create_from_upload(
            submission=submission,
            uploaded_file=SimpleUploadedFile("mail-extra.pdf", b"extra text"),
            document_type=SubmissionDocument.DocumentType.TEXT,
        )
        SubmissionDocument.create_from_upload(
            submission=submission,
            uploaded_file=SimpleUploadedFile("mail-visual.png", b"visual bytes"),
            document_type=SubmissionDocument.DocumentType.VISUAL,
        )

        sent = send_submission_receipt(submission)

        self.assertEqual(sent, 1)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].from_email, settings.SUBMISSION_EMAIL_FROM)
        self.assertEqual(mail.outbox[0].subject, settings.SUBMISSION_EMAIL_SUBJECT)
        self.assertIn("Thank you for submitting to Rage Zine.", mail.outbox[0].body)
        self.assertIn("Hi Artist Mail,", mail.outbox[0].body)
        self.assertIn('We received "Mail Piece"', mail.outbox[0].body)
        self.assertIn("Submissions are being reviewed", mail.outbox[0].body)
        self.assertIn("contact you at this email address to let you know the outcome", mail.outbox[0].body)
        self.assertIn(f"Submission reference: #{submission.pk}", mail.outbox[0].body)
        self.assertEqual(mail.outbox[0].to, ["artist.mail@example.com"])
        self.assertNotIn("Detailed description", mail.outbox[0].body)
        self.assertEqual(mail.outbox[0].attachments, [])
        self.assertNotIn("mail-piece.pdf", mail.outbox[0].body)
        self.assertNotIn("/downloads/", mail.outbox[0].body)
        self.assertEqual(len(mail.outbox[0].alternatives), 1)
        html, content_type = mail.outbox[0].alternatives[0]
        self.assertEqual(content_type, "text/html")
        self.assertIn("Thank you for sharing your work.", html)
        self.assertIn("Submissions are being reviewed", html)
        self.assertIn("contact you at this email address to let you know the outcome", html)
        self.assertIn("Artist Mail", html)
        self.assertNotIn("mail-piece.pdf", html)
        self.assertNotIn("/downloads/", html)

    @override_settings(
        EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
        SUBMISSION_EMAIL_FROM="editor@ragezine.test",
        SUBMISSION_EMAIL_BODY="Thanks <script>alert(1)</script>",
    )
    def test_receipt_html_escapes_submitted_text(self):
        submission = SimpleNamespace(
            pk=7,
            artist_name="<b>Artist</b>",
            title="<script>Piece</script>",
            email="artist@example.com",
        )

        send_submission_receipt(submission)

        html = mail.outbox[0].alternatives[0][0]
        self.assertNotIn("<script>", html)
        self.assertNotIn("<b>Artist</b>", html)
        self.assertIn("&lt;script&gt;Piece&lt;/script&gt;", html)
        self.assertIn("&lt;b&gt;Artist&lt;/b&gt;", html)
