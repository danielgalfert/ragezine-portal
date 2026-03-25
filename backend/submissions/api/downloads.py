from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from django.utils.text import get_valid_filename
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated

from submissions.models import Submission
from submissions.serializers import country_name
from submissions.utils import sanitize_single_line


EXPORT_FILENAME = "ragezine-submissions-complete.zip"
SUBMISSION_DETAILS_FILENAME = "submission_details.txt"
EXCEL_COLUMN_CONFIG = [
    ("id", "ID"),
    ("title", "Title"),
    ("year", "Year"),
    ("description", "Description"),
    ("submission_type", "Submission Type"),
    ("artist_name", "Artist Name"),
    ("pronouns", "Pronouns"),
    ("short_bio", "Short Bio"),
    ("socials", "Socials"),
    ("email", "Email"),
    ("country_origin", "Country Origin"),
    ("countries_residence", "Countries Residence"),
    ("language", "Language"),
    ("created_at", "Created At"),
]
WIDE_TEXT_FIELDS = {"description", "short_bio"}


def _safe_path_segment(value, fallback):
    cleaned = get_valid_filename(sanitize_single_line(value))
    return cleaned or fallback


def _submission_archive_name(submission):
    artist_segment = _safe_path_segment(
        submission.artist_name,
        f"submission-{submission.id}",
    )
    return f"ragezine-submission-{submission.id}-{artist_segment}.zip"


def _iter_submission_files(submission):
    if submission.text_file:
        yield submission.text_file

    for text in submission.texts.all():
        yield text.file

    for visual in submission.visuals.all():
        yield visual.image


def _format_submission_details(submission):
    countries_residence = ", ".join(
        country_name(code) for code in submission.countries_residence
    ) or "-"
    socials = submission.socials or "-"
    attached_files = [Path(stored_file.name).name for stored_file in _iter_submission_files(submission)]

    if not attached_files:
        attached_files = ["-"]

    lines = [
        f"Submission ID: {submission.id}",
        f"Title: {submission.title}",
        f"Submission Type: {submission.get_submission_type_display()} ({submission.submission_type})",
        f"Artist Name: {submission.artist_name}",
        f"Email: {submission.email}",
        f"Year: {submission.year or '-'}",
        f"Description: {submission.description or '-'}",
        f"Pronouns: {submission.pronouns}",
        f"Short Bio: {submission.short_bio}",
        f"Socials: {socials}",
        f"Country of Origin: {country_name(submission.country_origin)}",
        f"Countries of Residence: {countries_residence}",
        f"Language: {submission.language}",
        f"Allow Translation: {'yes' if submission.allow_translation else 'no'}",
        f"Created At: {submission.created_at.isoformat()}",
        f"Updated At: {submission.updated_at.isoformat()}",
        "Files:",
    ]
    lines.extend(f"- {filename}" for filename in attached_files)
    return "\n".join(lines) + "\n"


def _format_excel_field(submission, field_name):
    if field_name == "submission_type":
        return submission.get_submission_type_display()
    if field_name == "country_origin":
        return country_name(submission.country_origin)
    if field_name == "countries_residence":
        return ", ".join(country_name(code) for code in submission.countries_residence)
    if field_name == "text_file":
        return Path(submission.text_file.name).name if submission.text_file else ""
    if field_name == "created_at":
        return submission.created_at.replace(tzinfo=None)
    return getattr(submission, field_name)


def _resolve_excel_filename(submissions):
    volumes = sorted({submission.volume for submission in submissions if submission.volume is not None})
    if len(volumes) == 1:
        volume_number = volumes[0]
    else:
        volume_number = Submission._meta.get_field("volume").default
    return f"rage_submissions_volume{volume_number}.xlsx"


def _autosize_worksheet(worksheet):
    for column_index, (field_name, _) in enumerate(EXCEL_COLUMN_CONFIG, start=1):
        column_letter = worksheet.cell(row=1, column=column_index).column_letter
        if field_name in WIDE_TEXT_FIELDS:
            worksheet.column_dimensions[column_letter].width = 48
            continue

        max_length = 0
        for row in worksheet.iter_rows(
            min_row=1,
            max_row=worksheet.max_row,
            min_col=column_index,
            max_col=column_index,
        ):
            value = row[0].value
            if value is None:
                continue
            max_length = max(max_length, len(str(value)))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 36)


def _build_excel_workbook():
    submissions = list(Submission.objects.all())
    buffer = BytesIO()
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="top", wrap_text=False)
    long_text_alignment = Alignment(vertical="top", wrap_text=True)

    for submission_type in Submission.SubmissionType:
        worksheet = workbook.create_sheet(title=submission_type.label)
        worksheet.freeze_panes = "A2"
        worksheet.append([header for _, header in EXCEL_COLUMN_CONFIG])

        for column_index, (_, header) in enumerate(EXCEL_COLUMN_CONFIG, start=1):
            cell = worksheet.cell(row=1, column=column_index)
            cell.value = header
            cell.font = Font(bold=True)
            cell.alignment = header_alignment

        type_submissions = [
            submission for submission in submissions
            if submission.submission_type == submission_type.value
        ]
        for submission in type_submissions:
            worksheet.append([
                _format_excel_field(submission, field_name)
                for field_name, _ in EXCEL_COLUMN_CONFIG
            ])

        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for column_index, cell in enumerate(row, start=1):
                field_name = EXCEL_COLUMN_CONFIG[column_index - 1][0]
                cell.alignment = (
                    long_text_alignment if field_name in WIDE_TEXT_FIELDS else cell_alignment
                )
            if worksheet.max_row > 1:
                worksheet.row_dimensions[row[0].row].height = 60

        _autosize_worksheet(worksheet)

    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue(), _resolve_excel_filename(submissions)


def _build_complete_download_zip():
    buffer = BytesIO()
    submissions = Submission.objects.prefetch_related("texts", "visuals").all()

    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        for submission_type in Submission.SubmissionType:
            archive.writestr(f"{submission_type.value}/", b"")

        for submission in submissions:
            artist_segment = _safe_path_segment(
                submission.artist_name,
                f"submission-{submission.id}",
            )
            submission_root = f"{submission.submission_type}/{artist_segment}"
            archive.writestr(
                f"{submission_root}/{SUBMISSION_DETAILS_FILENAME}",
                _format_submission_details(submission),
            )

            used_names = set()
            for stored_file in _iter_submission_files(submission):
                file_name = _safe_path_segment(
                    Path(stored_file.name).name,
                    f"file-{len(used_names) + 1}",
                )
                while file_name in used_names:
                    stem = Path(file_name).stem
                    suffix = Path(file_name).suffix
                    file_name = f"{stem}-{len(used_names) + 1}{suffix}"
                used_names.add(file_name)

                stored_file.open("rb")
                try:
                    archive.writestr(
                        f"{submission_root}/{file_name}",
                        stored_file.read(),
                    )
                finally:
                    stored_file.close()

    buffer.seek(0)
    return buffer.getvalue()


def _write_submission_to_archive(archive, submission):
    artist_segment = _safe_path_segment(
        submission.artist_name,
        f"submission-{submission.id}",
    )
    submission_root = f"{submission.submission_type}/{artist_segment}"
    archive.writestr(
        f"{submission_root}/{SUBMISSION_DETAILS_FILENAME}",
        _format_submission_details(submission),
    )

    used_names = set()
    for stored_file in _iter_submission_files(submission):
        file_name = _safe_path_segment(
            Path(stored_file.name).name,
            f"file-{len(used_names) + 1}",
        )
        while file_name in used_names:
            stem = Path(file_name).stem
            suffix = Path(file_name).suffix
            file_name = f"{stem}-{len(used_names) + 1}{suffix}"
        used_names.add(file_name)

        stored_file.open("rb")
        try:
            archive.writestr(
                f"{submission_root}/{file_name}",
                stored_file.read(),
            )
        finally:
            stored_file.close()


def _build_submission_download_zip(submission):
    buffer = BytesIO()

    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        _write_submission_to_archive(archive, submission)

    buffer.seek(0)
    return buffer.getvalue(), _submission_archive_name(submission)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def complete_download(request):
    zip_bytes = _build_complete_download_zip()
    response = HttpResponse(zip_bytes, content_type="application/zip")
    response["Content-Disposition"] = f'attachment; filename="{EXPORT_FILENAME}"'
    response["Content-Length"] = str(len(zip_bytes))
    return response


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def submission_download(request, submission_id):
    submission = get_object_or_404(
        Submission.objects.prefetch_related("texts", "visuals"),
        pk=submission_id,
    )
    zip_bytes, export_filename = _build_submission_download_zip(submission)
    response = HttpResponse(zip_bytes, content_type="application/zip")
    response["Content-Disposition"] = f'attachment; filename="{export_filename}"'
    response["Content-Length"] = str(len(zip_bytes))
    return response


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def excel_download(request):
    workbook_bytes, export_filename = _build_excel_workbook()
    response = HttpResponse(
        workbook_bytes,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{export_filename}"'
    response["Content-Length"] = str(len(workbook_bytes))
    return response
